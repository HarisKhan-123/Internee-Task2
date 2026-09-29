"""
Resume text extraction utilities.
Supports .pdf and .docx files.
"""
import os
from pypdf import PdfReader
from docx import Document


class ResumeParsingError(Exception):
    """Raised when a resume file cannot be parsed."""
    pass


def extract_text_from_pdf(file_path: str) -> str:
    """Extract raw text from a PDF file."""
    try:
        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
        text = "\n".join(text_parts).strip()
        if not text:
            raise ResumeParsingError(
                "No extractable text found in the PDF. It may be a scanned "
                "image — try a text-based PDF or a DOCX file instead."
            )
        return text
    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ResumeParsingError(f"Failed to read PDF: {exc}") from exc


def extract_text_from_docx(file_path: str) -> str:
    """Extract raw text from a DOCX file, including tables."""
    try:
        doc = Document(file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also pull text out of any tables (skills/experience tables are common)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

        text = "\n".join(parts).strip()
        if not text:
            raise ResumeParsingError("No extractable text found in the DOCX file.")
        return text
    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ResumeParsingError(f"Failed to read DOCX: {exc}") from exc


def extract_resume_text(file_path: str) -> str:
    """Dispatch to the correct extractor based on file extension."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise ResumeParsingError(f"Unsupported file type: {ext}. Please upload a PDF or DOCX.")
